"""成绩：查询、分析、增删改、导入导出。"""
import os
from io import BytesIO

from fastapi import Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app.models import ImportHistory, Score, Student
from app.schemas import ScoreCreate, ScoreOut, ScoreUpdate
from app.routers.workbench._common import (
    dep,
    get_db,
    get_current_user,
    new_router,
    audit,
    student_name,
    active_student_id_query,
    batch_student_map,
    apply_student_class_filter,
    apply_teacher_student_filter,
    ensure_student_operable,
    is_any_admin,
    is_student_in_teacher_classes,
    get_teacher_class_ids,
    attach_student,
    serialize_list_with_students,
    to_dict,
    normalize_page,
)
from app.utils import clamp_score

router = new_router("成绩")


@router.get("/scores")
def list_scores(
    page: int = 1,
    page_size: int = 20,
    student_id: int | None = None,
    class_id: int | None = None,
    subject: str = "",
    user=Depends(get_current_user),
    db: Session = Depends(get_db),
):
    page, page_size = normalize_page(page, page_size)
    q = db.query(Score)
    # 排除退学学生
    q = q.filter(Score.student_id.in_(active_student_id_query(db)))
    # 教师只能查看自己负责班级的学生成绩
    q, denied = apply_teacher_student_filter(db, user, q, Score)
    if denied:
        return {"items": [], "total": 0}
    if class_id:
        q, denied = apply_student_class_filter(db, user, q, class_id, Score)
        if denied:
            return {"items": [], "total": 0}
    if student_id:
        if not is_any_admin(user) and not is_student_in_teacher_classes(db, user.id, student_id):
            return {"items": [], "total": 0}
        q = q.filter(Score.student_id == student_id)
    if subject:
        q = q.filter(Score.subject == subject)
    total = q.count()
    rows = q.order_by(Score.id.desc()).offset((page - 1) * page_size).limit(page_size).all()
    items = serialize_list_with_students(db, rows)
    return {"items": items, "total": total}


@router.get("/scores/analysis")
def score_analysis(
    class_id: int | None = None,
    student_id: int | None = None,
    subject: str = "",
    exam_name: str = "",
    user=Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """成绩分析：班级排名 + 学生成绩趋势 + 可选科目/考试维度。"""
    resp: dict = {"subjects": [], "exams": []}

    base_q = db.query(Score).filter(Score.student_id.in_(active_student_id_query(db)))
    base_q, denied = apply_teacher_student_filter(db, user, base_q, Score)
    if denied:
        return resp

    if class_id:
        base_q, denied = apply_student_class_filter(db, user, base_q, class_id, Score)
        if denied:
            return resp

    subjects = sorted({s for (s,) in base_q.with_entities(Score.subject).all()})
    exams = sorted({e for (e,) in base_q.with_entities(Score.exam_name).filter(Score.exam_name.isnot(None)).all()})
    resp["subjects"] = subjects
    resp["exams"] = exams

    if class_id:
        rank_q = db.query(Score).filter(
            Score.student_id.in_(active_student_id_query(db)),
            Score.student_id.in_(
                db.query(Student.id).filter(Student.class_id == class_id)
            ),
        )
        if subject:
            rank_q = rank_q.filter(Score.subject == subject)
        if exam_name:
            rank_q = rank_q.filter(Score.exam_name == exam_name)
        rows = rank_q.all()
        rank_map: dict = {}
        for r in rows:
            rank_map.setdefault(r.student_id, []).append(r)
        students = db.query(Student).filter(Student.id.in_(list(rank_map.keys()))).all()
        student_map = {s.id: s for s in students}
        ranking = []
        for sid, scores in rank_map.items():
            s = student_map.get(sid)
            if not s:
                continue
            vals = [x.score for x in scores]
            avg = round(sum(vals) / len(vals), 1)
            ranking.append({
                "student_id": sid,
                "name": s.name,
                "student_no": s.student_no,
                "score": avg,
                "count": len(vals),
            })
        ranking.sort(key=lambda x: x["score"], reverse=True)
        for i, item in enumerate(ranking):
            item["rank"] = i + 1
        resp["ranking"] = ranking
        resp["ranking_total"] = len(ranking)

    if student_id:
        if not is_any_admin(user) and not is_student_in_teacher_classes(db, user.id, student_id):
            raise HTTPException(status_code=403, detail="无权查看该学生成绩")
        sc_rows = (
            db.query(Score)
            .filter(Score.student_id == student_id)
            .order_by(Score.created_at)
            .all()
        )
        stu = db.get(Student, student_id)
        class_ids = [stu.class_id] if (stu and stu.class_id) else []
        class_avg_by_exam: dict = {}
        if class_ids:
            class_scores = (
                db.query(Score)
                .filter(Score.student_id.in_(active_student_id_query(db)))
                .filter(
                    Score.student_id.in_(
                        db.query(Student.id).filter(Student.class_id.in_(class_ids))
                    )
                )
                .all()
            )
            exam_group: dict = {}
            for r in class_scores:
                key = (r.subject, r.exam_name or "")
                exam_group.setdefault(key, []).append(r.score)
            for key, vals in exam_group.items():
                class_avg_by_exam[key] = round(sum(vals) / len(vals), 1)

        trend = []
        for r in sc_rows:
            key = (r.subject, r.exam_name or "")
            trend.append({
                "subject": r.subject,
                "score": r.score,
                "exam": r.exam_name or "",
                "date": str(r.created_at)[:10],
                "class_avg": class_avg_by_exam.get(key),
            })
        resp["trend"] = trend

        by_subject: dict = {}
        for r in sc_rows:
            by_subject.setdefault(r.subject, []).append({
                "score": r.score,
                "exam": r.exam_name or "",
                "date": str(r.created_at)[:10],
                "class_avg": class_avg_by_exam.get((r.subject, r.exam_name or "")),
            })
        resp["by_subject"] = by_subject

    return resp


@router.post("/scores", response_model=ScoreOut, status_code=201)
def create_score(payload: ScoreCreate, user=Depends(get_current_user), db: Session = Depends(get_db)):
    ensure_student_operable(db, payload.student_id)
    if not is_any_admin(user) and not is_student_in_teacher_classes(db, user.id, payload.student_id):
        raise HTTPException(status_code=403, detail="无权为该学生创建成绩")
    s = Score(
        student_id=payload.student_id,
        subject=payload.subject,
        score=payload.score,
        exam_name=payload.exam_name,
    )
    db.add(s)
    audit(db, user, "create_score", target=f"新增成绩-{student_name(db, s.student_id)}", student_id=s.student_id, detail=f"科目：{s.subject}；分数：{s.score}；考试：{s.exam_name or '日常测验'}")
    db.commit()
    db.refresh(s)
    return attach_student(db, to_dict(s), s.student_id)


@router.put("/scores/{score_id}", response_model=ScoreOut)
def update_score(score_id: int, payload: ScoreUpdate, user=Depends(get_current_user), db: Session = Depends(get_db)):
    s = db.get(Score, score_id)
    if not s:
        raise HTTPException(status_code=404, detail="记录不存在")
    ensure_student_operable(db, s.student_id)
    if not is_any_admin(user) and not is_student_in_teacher_classes(db, user.id, s.student_id):
        raise HTTPException(status_code=403, detail="无权修改该成绩")
    data = payload.model_dump(exclude_unset=True)
    for f in ("student_id", "subject", "score", "exam_name"):
        if f in data and data[f] is not None:
            if f == "student_id" and not is_any_admin(user):
                if not is_student_in_teacher_classes(db, user.id, data[f]):
                    raise HTTPException(status_code=403, detail="无权将成绩转移到该学生")
            if f == "student_id" and data[f] != s.student_id:
                ensure_student_operable(db, data[f])
            setattr(s, f, data[f])
    audit(db, user, "update_score", target=f"成绩#{score_id}-{student_name(db, s.student_id)}", student_id=s.student_id)
    db.commit()
    db.refresh(s)
    return attach_student(db, to_dict(s), s.student_id)


@router.delete("/scores/{score_id}")
def delete_score(score_id: int, user=Depends(get_current_user), db: Session = Depends(get_db)):
    s = db.get(Score, score_id)
    if not s:
        raise HTTPException(status_code=404, detail="记录不存在")
    ensure_student_operable(db, s.student_id)
    if not is_any_admin(user) and not is_student_in_teacher_classes(db, user.id, s.student_id):
        raise HTTPException(status_code=403, detail="无权删除该成绩")
    db.delete(s)
    audit(db, user, "delete_score", target=f"成绩#{score_id}-{student_name(db, s.student_id)}", student_id=s.student_id)
    db.commit()
    return {"ok": True}


def _empty_score_export():
    """无权限/无数据时返回空的成绩导出文件（与正常导出同构）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "成绩单"
    ws.append(["学号", "姓名", "科目", "成绩", "考试名称"])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=scores.xlsx"},
    )


@router.get("/scores/export")
def export_scores(student_id: int | None = None, class_id: int | None = None, user=Depends(get_current_user), db: Session = Depends(get_db)):
    q = db.query(Score)
    q = q.filter(Score.student_id.in_(active_student_id_query(db)))
    q, denied = apply_teacher_student_filter(db, user, q, Score)
    if denied:
        return _empty_score_export()
    if class_id:
        q, denied = apply_student_class_filter(db, user, q, class_id, Score)
        if denied:
            return _empty_score_export()
    if student_id:
        if not is_any_admin(user) and not is_student_in_teacher_classes(db, user.id, student_id):
            return _empty_score_export()
        q = q.filter(Score.student_id == student_id)
    rows = q.order_by(Score.student_id).all()
    stu_map = batch_student_map(db, [s.student_id for s in rows])
    wb = Workbook()
    ws = wb.active
    ws.title = "成绩单"
    ws.append(["学号", "姓名", "科目", "成绩", "考试名称"])
    for s in rows:
        stu = stu_map.get(s.student_id)
        ws.append([stu["no"] if stu else "", stu["name"] if stu else "", s.subject, s.score, s.exam_name or ""])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=scores.xlsx"},
    )


def _validate_score_row(row_data: dict, row_num: int) -> list[str]:
    """校验成绩导入行数据，返回错误列表。"""
    errors = []
    if not row_data.get("student_no", "").strip():
        errors.append(f"第{row_num}行：学号不能为空")
    if not row_data.get("subject", "").strip():
        errors.append(f"第{row_num}行：科目不能为空")
    try:
        float(row_data.get("score", 0))
    except (ValueError, TypeError):
        errors.append(f"第{row_num}行：成绩必须是数字")
    return errors


@router.get("/scores/template")
def download_score_template(_=Depends(dep)):
    """下载成绩导入模板。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "成绩导入模板"
    ws.append(["学号", "姓名", "科目", "成绩", "考试名称"])
    ws.append(["2024001", "张三", "语文", "85", "期中考试"])
    for col, w in enumerate([12, 10, 10, 8, 14], 1):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = w
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=score_template.xlsx"})


@router.post("/scores/import")
def import_scores(
    file: UploadFile = File(...),
    user=Depends(dep),
    db: Session = Depends(get_db),
):
    """批量导入成绩数据。"""
    ext = os.path.splitext(file.filename or "")[1].lower()
    if ext not in {".xlsx", ".xls"}:
        raise HTTPException(status_code=400, detail="仅支持 .xlsx / .xls 格式")

    contents = file.file.read()
    wb = load_workbook(BytesIO(contents))
    ws = wb.active

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    if not rows:
        raise HTTPException(status_code=400, detail="文件中没有数据行")

    all_errors = []
    success = 0
    total = 0

    if not is_any_admin(user):
        teacher_class_ids = get_teacher_class_ids(db, user.id)
        teacher_student_nos = set()
        if teacher_class_ids:
            for s in db.query(Student).filter(Student.class_id.in_(teacher_class_ids)).all():
                teacher_student_nos.add(s.student_no)
    else:
        teacher_student_nos = None

    student_map = {s.student_no: s for s in db.query(Student).all()}

    for row_num, row in enumerate(rows, start=2):
        if not any(row):
            continue
        total += 1
        data = {
            "student_no": str(row[0] or "").strip(),
            "name": str(row[1] or "").strip(),
            "subject": str(row[2] or "").strip(),
            "score": row[3],
            "exam_name": str(row[4] or "").strip(),
        }

        errors = _validate_score_row(data, row_num)
        if errors:
            all_errors.extend(errors)
            continue

        student = student_map.get(data["student_no"])
        if not student:
            all_errors.append(f"第{row_num}行：学号 {data['student_no']} 不存在")
            continue

        if student.is_dropped_out:
            all_errors.append(f"第{row_num}行：学生「{data['name']}」已退学，无法导入成绩")
            continue
        if student.class_id:
            cls = db.get(Classroom, student.class_id)
            if cls and cls.is_graduated:
                all_errors.append(f"第{row_num}行：学生「{data['name']}」所在班级已毕业，无法导入成绩")
                continue

        if teacher_student_nos is not None and data["student_no"] not in teacher_student_nos:
            all_errors.append(f"第{row_num}行：教师只能导入自己班级学生「{data['name']}」的成绩")
            continue

        try:
            db.add(Score(
                student_id=student.id,
                subject=data["subject"],
                score=float(data["score"]),
                exam_name=data["exam_name"],
            ))
            success += 1
        except Exception as e:
            all_errors.append(f"第{row_num}行：导入失败 - {str(e)}")

    db.commit()
    audit(db, user, "import_scores", target=f"{file.filename or ''} 成功{success}条")

    db.add(ImportHistory(
        import_type="score",
        filename=file.filename or "",
        total_rows=total,
        success_rows=success,
        error_rows=len(all_errors),
        errors=json.dumps(all_errors[:100], ensure_ascii=False),
        user_id=user.id,
    ))
    db.commit()

    return {"success": success, "total": total, "errors": all_errors[:50]}
